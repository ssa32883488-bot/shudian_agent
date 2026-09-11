"""错题本 CRUD + Word 导出。"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Optional

from docx import Document
from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models_ext import Mistake


def list_mistakes(db: Session, student_id: int) -> list[Mistake]:
    return list(
        db.scalars(
            select(Mistake)
            .where(Mistake.student_id == student_id, Mistake.deleted_at.is_(None))
            .order_by(Mistake.added_at.desc())
        ).all()
    )


def add_mistake(
    db: Session,
    *,
    student_id: int,
    question_snapshot: str,
    answer_snapshot: Optional[str] = None,
    reason: Optional[str] = None,
    question_id: Optional[int] = None,
    add_source: str = "manual",
    knowledge_tags: Optional[dict] = None,
) -> Mistake:
    m = Mistake(
        student_id=student_id,
        question_id=question_id,
        question_snapshot=question_snapshot,
        answer_snapshot=answer_snapshot,
        reason=reason,
        add_source=add_source,
        knowledge_tags=knowledge_tags,
    )
    db.add(m)
    db.commit()
    db.refresh(m)
    return m


def update_mistake(
    db: Session, mistake_id: int, student_id: int, **fields
) -> Mistake:
    m = db.get(Mistake, mistake_id)
    if not m or m.student_id != student_id or m.deleted_at:
        raise HTTPException(status_code=404, detail="错题不存在")
    for k, v in fields.items():
        if v is not None and hasattr(m, k):
            setattr(m, k, v)
    db.commit()
    db.refresh(m)
    return m


def delete_mistake(db: Session, mistake_id: int, student_id: int) -> None:
    m = db.get(Mistake, mistake_id)
    if not m or m.student_id != student_id:
        raise HTTPException(status_code=404, detail="错题不存在")
    m.deleted_at = datetime.now(timezone.utc)
    db.commit()


def import_from_exam(
    db: Session, student_id: int, items: list[dict]
) -> list[Mistake]:
    out: list[Mistake] = []
    for it in items:
        source = str(it.get("add_source") or "manual_import").strip() or "manual_import"
        out.append(
            add_mistake(
                db,
                student_id=student_id,
                question_snapshot=it["question"],
                answer_snapshot=it.get("answer"),
                reason=it.get("reason", "批量导入"),
                question_id=it.get("question_id"),
                add_source=source,
                knowledge_tags=it.get("knowledge_tags"),
            )
        )
    return out


def export_word(db: Session, student_id: int) -> bytes:
    mistakes = list_mistakes(db, student_id)
    doc = Document()
    doc.add_heading("数电错题本", level=1)
    for i, m in enumerate(mistakes, 1):
        doc.add_heading(f"第 {i} 题", level=2)
        doc.add_paragraph(m.question_snapshot)
        if m.answer_snapshot:
            doc.add_paragraph(f"参考答案：{m.answer_snapshot}")
        if m.reason:
            doc.add_paragraph(f"复盘：{m.reason}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
