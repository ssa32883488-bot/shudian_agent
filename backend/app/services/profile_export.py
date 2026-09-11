"""个人学情 Word 导出。"""

from __future__ import annotations

import io
from typing import Any

from docx import Document


def export_student_profile_word(payload: dict[str, Any]) -> bytes:
    doc = Document()
    name = payload.get("nickname") or "同学"
    doc.add_heading(f"个人学情报告 · {name}", level=1)
    if payload.get("class_code"):
        doc.add_paragraph(f"班级码：{payload['class_code']}")
    if payload.get("updated_at"):
        doc.add_paragraph(f"画像更新：{payload['updated_at']}")
    if payload.get("last_rebuild_at"):
        doc.add_paragraph(f"最近重建：{payload['last_rebuild_at']}")

    stats = payload.get("practice_stats") or {}
    doc.add_heading("练习概况", level=2)
    doc.add_paragraph(f"累计刷题：{stats.get('total_solves', 0)}")
    acc = stats.get("accuracy")
    if acc is not None:
        doc.add_paragraph(f"最近正确率：{round(float(acc) * 100)}%")
    last = stats.get("last_session") or {}
    if last:
        doc.add_paragraph(
            f"最近一场：对 {last.get('correct', '—')}/{last.get('total', '—')} · "
            f"{last.get('at', '')}"
        )
        missed = last.get("missed_tags") or []
        if missed:
            doc.add_paragraph("本场错点：" + "、".join(str(x) for x in missed))

    mastery = payload.get("mastery") or {}
    if mastery:
        doc.add_heading("知识点掌握度", level=2)
        for k, v in sorted(mastery.items(), key=lambda kv: float(kv[1])):
            doc.add_paragraph(f"{k}：{round(float(v) * 100)}%")

    weak = payload.get("weak_points") or []
    if weak:
        doc.add_heading("薄弱知识点", level=2)
        for w in weak:
            label = w if isinstance(w, str) else (w.get("concept") if isinstance(w, dict) else str(w))
            doc.add_paragraph(str(label))

    mistake_stats = payload.get("mistake_stats") or stats.get("mistake_stats") or {}
    top_tags = mistake_stats.get("top_tags") or []
    if top_tags:
        doc.add_heading("错题标签分布", level=2)
        for item in top_tags:
            if isinstance(item, dict):
                doc.add_paragraph(f"{item.get('name')} · {item.get('count')} 次")
            else:
                doc.add_paragraph(str(item))
    elif mistake_stats.get("count"):
        doc.add_heading("错题", level=2)
        doc.add_paragraph(f"错题本共 {mistake_stats.get('count')} 道")

    topics = payload.get("recent_topics") or []
    if topics:
        doc.add_heading("近期关注话题", level=2)
        doc.add_paragraph("、".join(str(t) for t in topics))

    freq = payload.get("topic_freq") or {}
    if freq:
        doc.add_heading("话题频次", level=2)
        for k, v in sorted(freq.items(), key=lambda kv: -int(kv[1] or 0))[:20]:
            doc.add_paragraph(f"{k}：{v}")

    recs = payload.get("recommendations") or []
    if recs:
        doc.add_heading("学习建议", level=2)
        for r in recs:
            if isinstance(r, dict):
                doc.add_paragraph(str(r.get("label") or r))
            else:
                doc.add_paragraph(str(r))

    note = payload.get("privacy_note")
    if note:
        doc.add_paragraph(str(note))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
