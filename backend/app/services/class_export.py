"""班级学情 Word 导出。"""

from __future__ import annotations

import io
from typing import Any

from docx import Document


def export_class_learning_word(payload: dict[str, Any]) -> bytes:
    doc = Document()
    code = payload.get("class_code") or ""
    doc.add_heading(f"班级学情报告 · {code}", level=1)
    doc.add_paragraph(f"学生人数：{payload.get('n_students', 0)}")
    if payload.get("computed_at"):
        doc.add_paragraph(f"聚合时间：{payload['computed_at']}")

    mastery = payload.get("mastery_avg") or {}
    if mastery:
        doc.add_heading("知识点掌握度均值", level=2)
        for k, v in mastery.items():
            doc.add_paragraph(f"{k}：{round(float(v) * 100)}%")

    weak = payload.get("weak_top") or []
    if weak:
        doc.add_heading("班级薄弱点 Top", level=2)
        for item in weak:
            doc.add_paragraph(f"{item.get('name')} · {item.get('count')} 人")

    mistakes = payload.get("mistake_top") or []
    if mistakes:
        doc.add_heading("错题标签分布", level=2)
        for item in mistakes:
            doc.add_paragraph(f"{item.get('name')} · {item.get('count')} 次")

    practice = payload.get("practice") or {}
    if practice:
        doc.add_heading("练习概况", level=2)
        acc = practice.get("avg_accuracy")
        acc_s = f"{round(float(acc) * 100)}%" if acc is not None else "—"
        doc.add_paragraph(
            f"已提交场次 {practice.get('n_sessions', 0)} · "
            f"参与学生 {practice.get('n_students', 0)} · "
            f"平均正确率 {acc_s}"
        )

    roster = payload.get("roster") or []
    if roster:
        doc.add_heading("学生名册", level=2)
        for s in roster:
            weak_s = "、".join(s.get("weak_points") or []) or "暂无"
            doc.add_paragraph(
                f"{s.get('nickname')} · 掌握均值 "
                f"{round(float(s.get('mastery_avg') or 0) * 100)}% · 薄弱：{weak_s}"
            )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
