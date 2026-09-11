# -*- coding: utf-8 -*-
"""重写《助学智能体技术设计说明书》——专业技术文档叙事。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "助学智能体技术设计说明书.docx"


def set_run_font(run, *, east_asia="宋体", ascii_font="Times New Roman", size_pt=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = ascii_font
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east_asia)
    rFonts.set(qn("w:cs"), ascii_font)


def fmt(p, *, first_line=True, space_after=6, line_spacing=1.5, align=None, space_before=0):
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    pf.first_line_indent = Cm(0.74) if first_line else Cm(0)
    if align is not None:
        p.alignment = align


def body(doc, text, *, first_line=True):
    p = doc.add_paragraph()
    fmt(p, first_line=first_line)
    set_run_font(p.add_run(text), east_asia="宋体", size_pt=12)


def h(doc, text, level=1):
    p = doc.add_paragraph()
    fmt(p, first_line=False, space_after=10, line_spacing=1.3, space_before=16 if level == 1 else 11)
    if level == 1:
        set_run_font(p.add_run(text), east_asia="黑体", ascii_font="Arial", size_pt=16, bold=True)
    elif level == 2:
        set_run_font(p.add_run(text), east_asia="黑体", ascii_font="Arial", size_pt=14, bold=True)
    else:
        set_run_font(p.add_run(text), east_asia="楷体", ascii_font="Arial", size_pt=12, bold=True)


def bullet(doc, text, *, lead=None):
    p = doc.add_paragraph()
    fmt(p, first_line=False, space_after=4, line_spacing=1.35)
    p.paragraph_format.left_indent = Cm(0.74)
    if lead:
        set_run_font(p.add_run("• " + lead), east_asia="宋体", size_pt=12, bold=True)
        set_run_font(p.add_run(text), east_asia="宋体", size_pt=12)
    else:
        set_run_font(p.add_run("• " + text), east_asia="宋体", size_pt=12)


def code(doc, text):
    for line in text.strip("\n").splitlines():
        p = doc.add_paragraph()
        fmt(p, first_line=False, space_after=0, line_spacing=1.15)
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(line if line else " ")
        set_run_font(run, east_asia="Consolas", ascii_font="Consolas", size_pt=9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def cell(c, text, *, bold=False, center=False, size=10.5, east="宋体"):
    c.text = ""
    p = c.paragraphs[0]
    fmt(p, first_line=False, space_after=2, line_spacing=1.2)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(text), east_asia=east, size_pt=size, bold=bold)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, hd in enumerate(headers):
        cell(t.rows[0].cells[i], hd, bold=True, center=True, east="黑体")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell(t.rows[ri + 1].cells[ci], val, size=10)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(2.8)

    p = doc.add_paragraph()
    fmt(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    set_run_font(p.add_run("学灵伴 · 助学智能体技术设计说明书"), east_asia="黑体", ascii_font="Arial", size_pt=20, bold=True)

    p = doc.add_paragraph()
    fmt(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    set_run_font(
        p.add_run("运行时分层 · 任务编排 · 工具注册 · 绘图工作流"),
        east_asia="楷体",
        size_pt=12,
        color=RGBColor(0x44, 0x44, 0x44),
    )

    h(doc, "1. 设计目标", 1)
    body(
        doc,
        "助学智能体是学灵伴的业务中枢。设计目标是在数字电子技术学科约束下，"
        "将多模态题面理解、任务规划、权威知识检索、工具增强求解与专业出图组织为一条可控制、可观测的执行链路。"
        "系统强调职责分离：大模型负责意图理解与综合表述，检索、绘图与记忆写入由注册工具完成；"
        "长任务通过计划与人在回路获得可控性，短任务保持低交互摩擦。",
    )

    h(doc, "2. 运行时分层", 1)
    table(
        doc,
        ["逻辑层", "职责", "不做"],
        [
            ["感知层", "文本、图像识读、课件正文归一为题面", "工具选择、最终作答"],
            ["规划与路由层", "意图判定、短/长路径决策、任务切片", "具体渲染与向量检索细节"],
            ["编排执行层", "解题图调度与 ReAct 工具循环", "擅自将检索结果升格为权威"],
            ["护栏与可观测", "配额、超时、熔断、轨迹与步骤事件", "改写学科结论正文"],
        ],
        widths=[3.2, 6.5, 5.5],
    )
    body(
        doc,
        "记忆、工具与交付策略作为横切能力贯穿各层：短期会话保证多轮连贯，长期记忆与学情画像按需注入；"
        "交付层区分权威解答与 AI 参考解答。",
    )

    h(doc, "3. 控制拓扑", 1)
    body(
        doc,
        "系统采用双路径汇入单一执行内核的拓扑。外层负责任务形态控制，内层负责权威检索与工具增强作答；"
        "最终求解共享同一导师人格与同一工具注册表，避免规划与执行漂移。",
    )
    code(
        doc,
        """路径 A  任务控制面
  感知 → 意图路由
    · 短路径 → 直接进入解题编排
    · 长路径 → 任务清单确认 → 按原子任务进入解题编排

路径 B  解题编排面
  归一 → 识读 → 决策 → 题库预取
    · 命中 → 权威交付
    · 未命中 → 知识锚定 → ReAct 求解 → 校验 → 回流裁决

执行内核  聊天模型 + 系统规约 + 工具注册表 + 图执行器 + 断路器""",
    )

    h(doc, "4. 感知与输入通道", 1)
    table(
        doc,
        ["通道", "处理策略"],
        [
            ["纯文本", "直接作为题面进入路由与编排"],
            ["图像", "多模态识读抽取题干；下游默认题面已可读，禁止否认可见性"],
            ["docx / pdf / ppt", "抽取正文后参与意图路由与任务切片"],
        ],
        widths=[4.0, 11.2],
    )
    body(
        doc,
        "感知层仅完成内容归一与规模估计。复合请求的原子化由路由层在长路径下完成；"
        "单回合执行规模设有上限，超额部分通过续跑推进。",
    )

    h(doc, "5. 意图路由与双模式控制", 1)
    h(doc, "5.1 设计动机", 2)
    body(
        doc,
        "单题答疑与整卷解析对时延、上下文预算与失败半径的要求不同。"
        "因此在进入解题内核前引入路由层，将请求映射为短路径或长路径，而不是用同一种交互形态覆盖全部场景。",
    )
    h(doc, "5.2 模式定义", 2)
    table(
        doc,
        ["模式", "适用情境", "控制策略"],
        [
            ["react_short", "寒暄、单点讲解、单题及少量原子任务", "不强制确认，直接进入解题编排"],
            ["plan_execute", "多题、整卷、长文档、大纲类复合任务", "任务清单经人在回路确认后串行执行，支持续跑"],
        ],
        widths=[3.2, 5.8, 6.2],
    )
    h(doc, "5.3 路由契约", 2)
    body(
        doc,
        "路由模型的输出契约为结构化结果：任务类型、执行模式、原子任务清单、容量裁决与安全阻断标记。"
        "路由阶段禁止解题、讲课或出图。解析失败时降级为整段单一任务，优先保证可用性。"
        "安全策略对越界请求执行阻断，阻止其进入工具链。",
    )

    h(doc, "6. 解题编排图", 1)
    body(
        doc,
        "解题过程以有状态图表达。每个节点承担单一职责，通过条件边实现权威短路与自研路径分支。",
    )
    code(
        doc,
        """normalize → ocr_or_text → decide → scope_task → bank_prefetch
    ├─ 命中 → bank_done → 权威交付
    └─ 未命中 → kg_ground → ai_solve → validate → 回流 / 跳过回流""",
    )
    table(
        doc,
        ["节点", "职责"],
        [
            ["normalize / ocr_or_text", "统一题面与来源，完成图像识读合并"],
            ["decide", "产出轻量意图与步骤标记，不直接撰写最终答卷"],
            ["bank_prefetch", "解题意图下预取题库；命中则权威短路"],
            ["kg_ground", "按需注入知识图谱上下文"],
            ["ai_solve", "进入导师内核的工具增强求解"],
            ["validate / reflow", "校验交付质量；裁决是否进入内容回流"],
        ],
        widths=[4.5, 10.7],
    )
    body(
        doc,
        "题库预取位于自研循环之前，避免与工具内重复搜库。"
        "交付语义区分权威解答与 AI 参考解答；适合沉淀的未命中题目进入回流队列，经管理审核后入库。",
    )

    h(doc, "7. 导师内核与护栏", 1)
    body(
        doc,
        "导师内核由聊天模型、系统规约、工具注册表与执行器组成。"
        "系统规约定义数电助教人格、场景边界、题库使用约定、工具清单与绘图占位规则；"
        "用户侧注入题面、预检索说明、记忆摘要与图谱上下文等原始材料。",
    )
    table(
        doc,
        ["护栏维度", "策略"],
        [
            ["迭代深度", "限制工具循环步数"],
            ["时间预算", "墙钟超时与空闲超时"],
            ["工具配额", "绘图、教材检索、题库搜索分别计数熔断"],
            ["重复动作", "同名同参连续重复触发断路"],
            ["输出规模", "限制最大输出，抑制生成坍缩"],
        ],
        widths=[3.5, 11.7],
    )
    body(
        doc,
        "执行过程以流式状态事件向客户端呈现步骤进展；完整轨迹可供管理侧运维复盘。"
        "护栏默认不改写学科结论，只决定是否停止或降级交付。",
    )

    h(doc, "8. 工具注册表", 1)
    body(
        doc,
        "下列工具构成导师智能体的全部外部能力面。绘图能力收敛为单一门面，降低工具选择歧义。",
    )
    table(
        doc,
        ["工具", "能力说明"],
        [
            ["draw_with_workflow", "按图类生成并校验中间表示，交确定性引擎输出 SVG；答案以占位符绑定产物"],
            ["search_question_bank", "题库原题检索；仅高置信且通过同题校验时返回权威内容"],
            ["retrieve_multi_rerank", "教材多路召回与精排；插图以可溯源占位进入上下文"],
            ["graph_query", "查询前置路径、关联链路与相关练习线索"],
            ["read_profile", "读取学情聚合画像"],
            ["memory_search", "检索长期记忆以支持个性化"],
            ["memory_upsert", "写入偏好、易错点与笔记等记忆"],
            ["memory_forget", "按用户请求软删除记忆"],
        ],
        widths=[4.2, 11.0],
    )

    h(doc, "9. 绘图工作流", 1)
    h(doc, "9.1 定位", 2)
    body(
        doc,
        "绘图是解题编排的附属能力：正确性依赖结构化中间表示与校验重试，渲染交给确定性引擎。"
        "工作流不面向独立 CAD 产品形态，而服务于教材级图示表达。",
    )
    h(doc, "9.2 处理流水线", 2)
    code(
        doc,
        """调用 draw_with_workflow(brief, kind?, slots?, script?, retries)
  → 解析图类（显式参数 / 槽位 / 题意关键词）
  → 获取 IR（显式脚本 → 完备槽位 → 确定性回退 → Skill 引导生成）
  → 结构校验 → 渲染 SVG
  → 失败则反馈错误并在有限次数内重试
  → 成功则返回产物元数据；答案使用 DRAW 占位绑定 URL""",
    )
    h(doc, "9.3 图类目录", 2)
    table(
        doc,
        ["图类", "用途", "路径要点"],
        [
            ["logic_dag", "门级逻辑与表达式", "逻辑 IR → netlistsvg"],
            ["timing_wave", "时序波形", "波形 IR → 专用渲染"],
            ["state_machine", "状态转换", "状态 IR → 专用渲染"],
            ["truth_table / kmap / seven_seg", "真值表、卡诺图、七段", "参数化绘制"],
            ["char_curve", "传输特性曲线", "科学绘图"],
            ["msi_design", "中规模器件设计电路", "逻辑先行校验 → 电路渲染"],
        ],
        widths=[4.2, 5.0, 6.0],
    )
    body(
        doc,
        "芯片搭接、置零/加载、555、DAC、模 N 计数等走设计电路图类；门级表达式与基本门实现走逻辑符号图类。"
        "设计电路扩展点位于器件引脚库。会话级绘图调用受配额约束；题面要求出图而产物缺失时，收尾阶段可再次进入同一工作流补全。",
    )

    h(doc, "10. 记忆体系", 1)
    table(
        doc,
        ["层次", "机制"],
        [
            ["短期会话", "服务端消息为权威历史，支撑多轮对话"],
            ["长期记忆", "经 memory_* 工具读写，跨会话复用"],
            ["学情画像", "批处理聚合掌握度与薄弱点，供工具与练习消费"],
        ],
        widths=[3.2, 12.0],
    )

    h(doc, "11. 典型运行时序", 1)
    bullet(doc, "感知完成题面归一，路由判定任务模式。")
    bullet(doc, "长路径生成任务清单并完成人在回路确认；短路径直接进入解题编排。")
    bullet(doc, "题库预取命中则权威交付；否则知识锚定并进入工具增强求解与出图。")
    bullet(doc, "校验后按策略回流；超额任务通过续跑继续推进。")
    bullet(doc, "会话与学情沉淀；回流内容经审核进入题库，提升后续检索命中质量。")

    h(doc, "12. 小结", 1)
    body(
        doc,
        "助学智能体通过分层运行时、双模式任务控制、权威优先的解题编排、收敛的工具注册表与教材级绘图工作流，"
        "将学科垂类模型能力组织为可控制、可审计、可经营的助学执行系统。"
        "其价值不在于单次生成的长度，而在于把「读懂—规划—检索—作答—出图—沉淀」固化为稳定的工程闭环。",
    )

    p = doc.add_paragraph()
    fmt(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18)
    set_run_font(p.add_run("— 全文完 —"), east_asia="宋体", size_pt=10.5, color=RGBColor(0x66, 0x66, 0x66))

    doc.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    build()
