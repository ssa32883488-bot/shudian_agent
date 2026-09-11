# -*- coding: utf-8 -*-
"""重写《学灵伴 · 系统技术架构说明书》——专业架构文档叙事，无「已实现」等工程备忘措辞。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "系统技术架构说明书.docx"


def set_run_font(run, *, east_asia="宋体", ascii_font="Times New Roman", size_pt=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = ascii_font
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east_asia)
    rFonts.set(qn("w:cs"), ascii_font)


def fmt(p, *, first_line=True, space_after=6, line_spacing=1.5, align=None, space_before=0):
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    pf.first_line_indent = Cm(0.74) if first_line else Cm(0)
    if align is not None:
        p.alignment = align


def body(doc, text, *, first_line=True):
    p = doc.add_paragraph()
    fmt(p, first_line=first_line)
    set_run_font(p.add_run(text), east_asia="宋体", size_pt=12)
    return p


def h(doc, text, level=1):
    p = doc.add_paragraph()
    fmt(p, first_line=False, space_after=10, line_spacing=1.3, space_before=16 if level == 1 else 11)
    if level == 1:
        set_run_font(p.add_run(text), east_asia="黑体", ascii_font="Arial", size_pt=16, bold=True)
    elif level == 2:
        set_run_font(p.add_run(text), east_asia="黑体", ascii_font="Arial", size_pt=14, bold=True)
    else:
        set_run_font(p.add_run(text), east_asia="楷体", ascii_font="Arial", size_pt=12, bold=True)
    return p


def bullet(doc, text, *, lead=None):
    p = doc.add_paragraph()
    fmt(p, first_line=False, space_after=4, line_spacing=1.35)
    p.paragraph_format.left_indent = Cm(0.74)
    if lead:
        set_run_font(p.add_run("• " + lead), east_asia="宋体", size_pt=12, bold=True)
        set_run_font(p.add_run(text), east_asia="宋体", size_pt=12)
    else:
        set_run_font(p.add_run("• " + text), east_asia="宋体", size_pt=12)


def code(doc, text):
    for line in text.strip("\n").splitlines():
        p = doc.add_paragraph()
        fmt(p, first_line=False, space_after=0, line_spacing=1.15)
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(line if line else " ")
        set_run_font(run, east_asia="Consolas", ascii_font="Consolas", size_pt=9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def cell(c, text, *, bold=False, center=False, size=10.5, east="宋体"):
    c.text = ""
    p = c.paragraphs[0]
    fmt(p, first_line=False, space_after=2, line_spacing=1.2)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(text), east_asia=east, size_pt=size, bold=bold)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, hd in enumerate(headers):
        cell(t.rows[0].cells[i], hd, bold=True, center=True, east="黑体")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell(t.rows[ri + 1].cells[ci], val, size=10)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(2.8)

    # Cover
    p = doc.add_paragraph()
    fmt(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    set_run_font(p.add_run("学灵伴"), east_asia="黑体", ascii_font="Arial", size_pt=26, bold=True)

    p = doc.add_paragraph()
    fmt(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    set_run_font(p.add_run("系统技术架构说明书"), east_asia="黑体", ascii_font="Arial", size_pt=18, bold=True)

    p = doc.add_paragraph()
    fmt(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    set_run_font(p.add_run("数字电子技术 · 单学科教育智能体"), east_asia="楷体", size_pt=12)

    p = doc.add_paragraph()
    fmt(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    set_run_font(
        p.add_run("文档类型：Architecture Description　受众：技术评审 / 答辩评委　工程代号：shudian_agent"),
        east_asia="宋体",
        size_pt=10,
        color=RGBColor(0x55, 0x55, 0x55),
    )

    # —— 1 引言与目标 ——
    h(doc, "1. 引言与目标", 1)
    h(doc, "1.1 系统定位", 2)
    body(
        doc,
        "学灵伴是面向高等教育《数字电子技术》课程的学科垂类智能应用。"
        "系统以助学智能体为业务中枢，在教材知识、题库与专业绘图工具的约束下，"
        "为学生提供伴随式问答、解题讲解、学情驱动练习与错题管理；"
        "同时通过管理控制台完成题库经营、内容回流审核与班级学情分析。"
        "架构设计的核心命题不是「接入一个大模型」，而是将学科规范、可信知识与可编排工作流固化为可运营的助学系统。",
    )

    h(doc, "1.2 质量目标", 2)
    table(
        doc,
        ["优先级", "质量目标", "架构含义"],
        [
            ["高", "学科可信", "题库权威路径优先；AI 生成与权威内容分轨交付；关键结论可溯源"],
            ["高", "复杂任务可控", "长任务计划化与人在回路确认；工具调用受配额与超时约束"],
            ["高", "专业表达完备", "数电图类经结构化中间表示与确定性渲染，而非自由矢量涂鸦"],
            ["中", "可观测可运维", "执行轨迹可追踪；熔断与降级路径明确"],
            ["中", "可部署可复现", "容器化交付；知识灌库与运行时依赖可说明"],
        ],
        widths=[2.2, 3.5, 9.5],
    )

    h(doc, "1.3 干系人视角", 2)
    table(
        doc,
        ["角色", "主要诉求", "系统响应"],
        [
            ["学生", "问得懂、解得对、图画得规范", "双模式对话、权威/参考分级、绘图工作流"],
            ["教师管理员", "题库可经营、班级学情可见", "管理控制台：题库、回流、班级、学情导出"],
            ["技术评审", "架构清晰、可演进", "本文：上下文、构件、运行时、决策与展望"],
        ],
        widths=[3.0, 5.5, 6.7],
    )

    # —— 2 约束与边界 ——
    h(doc, "2. 约束与阶段聚焦", 1)
    h(doc, "2.1 阶段聚焦", 2)
    body(
        doc,
        "本阶段优先把「助学中枢」做深：学生问答—解题—出图—学情练习闭环，以及管理侧题库经营与班级学情分析。"
        "在此底座之上，考务协作、按章护航教学等方向作为产品下一阶段的自然延伸，详见第 9 章展望。",
    )
    h(doc, "2.2 技术约束", 2)
    bullet(doc, "单学科纵深：交互与知识工程限定于数字电子技术，拒绝做成跨学科通用助手。")
    bullet(doc, "模型与工具分离：推理与表述由大模型承担；检索、绘图、记忆写入由注册工具执行。")
    bullet(doc, "信任不可混淆：题库命中结果与模型自研结果在交付语义上严格区分。")
    bullet(doc, "密钥治理：支持客户端自备密钥（BYOK）；用户密钥不上服务器持久化。")

    # —— 3 解决方案策略 ——
    h(doc, "3. 解决方案策略", 1)
    body(
        doc,
        "策略上采用「中枢智能体 + 知识底座 + 经营面」三元结构。"
        "中枢智能体负责意图理解、任务编排与工具调用；知识底座（教材向量、题库、图谱）提供可引用证据；"
        "经营面（管理控制台）使未命中题目可回流、班级学情可聚合，从而形成内容自生长闭环。",
    )
    body(
        doc,
        "编排层借鉴业界 Agent 工程的主流范式：以有状态图描述解题流水线，以 ReAct 循环承载开放式工具使用，"
        "并以路由层区分短路径与长路径，避免将所有请求压入同一种交互形态。"
        "外层负责任务切片与人机确认闸门，内层负责权威检索与工具增强作答，两者汇入同一导师人格与同一工具注册表，"
        "保证行为一致性，消除「规划一套、执行另一套」的双逻辑风险。",
    )

    # —— 4 上下文 ——
    h(doc, "4. 上下文视图", 1)
    h(doc, "4.1 业务上下文", 2)
    code(
        doc,
        """学生 ──问答/练习/错题──► 学灵伴学生端 ──API──► 助学智能体运行时
教师管理员 ──题库/回流/班级──► 管理控制台 ──────────┘
                                    │
                    ┌───────────────┼───────────────┐
                    ▼               ▼               ▼
              教材/题库向量      知识图谱        大模型与OCR服务
              （Chroma等）     （校本图谱）      （OpenAI兼容网关）""",
    )
    h(doc, "4.2 逻辑分层", 2)
    code(
        doc,
        """展示层   Vue 3 应用（学生端 · 管理控制台）
接入层   Nginx（静态资源与 API 反向代理）
业务层   FastAPI · LangGraph 编排 · 工具网关 · 业务服务
数据层   关系库 · 向量库 · 图谱数据 · 媒体产物 · 会话与画像""",
    )

    # —— 5 构件 ——
    h(doc, "5. 构建块视图", 1)
    h(doc, "5.1 容器级构件", 2)
    table(
        doc,
        ["构件", "职责", "关键技术"],
        [
            ["学生 Web 应用", "对话、练习、错题、学情、图谱可视化", "Vue 3 / Pinia / KaTeX / ECharts"],
            ["管理控制台", "题库、回流审核、班级、学情导出、轨迹查阅", "同前端应用，角色路由隔离"],
            ["API 与编排服务", "鉴权、计划路由、解题图、ReAct、回流与学情", "FastAPI / LangGraph / LangChain"],
            ["知识与检索子系统", "教材切块检索、题库召回、图谱查询", "Chroma / BGE Embedding·Rerank"],
            ["绘图工作流", "图类路由、IR 校验、确定性渲染", "Skill + Runner（netlistsvg 等）"],
            ["持久化与媒体", "账号班级、题库会话、画像、SVG 产物", "SQLAlchemy · 本地/对象存储"],
        ],
        widths=[3.5, 6.0, 5.7],
    )

    h(doc, "5.2 智能体运行时（系统中枢）", 2)
    body(
        doc,
        "智能体运行时划分为感知、规划与路由、编排执行、护栏与可观测四个逻辑层。"
        "感知层将文本、图像识读结果与课件正文归一为题面；规划与路由层判定任务形态并产出执行计划；"
        "编排执行层按图节点或 ReAct 循环调度工具；护栏层以迭代上限、墙钟超时、工具配额与重复动作熔断约束运行时行为，"
        "并通过流式状态事件与执行轨迹支撑可观测性。",
    )

    h(doc, "5.3 工具注册表", 2)
    body(
        doc,
        "导师智能体通过统一工具注册表获得外部能力。工具面保持收敛，避免同义多入口导致选择漂移。",
    )
    table(
        doc,
        ["工具", "能力边界"],
        [
            ["draw_with_workflow", "唯一绘图入口；按图类生成并校验中间表示，再交确定性引擎输出 SVG"],
            ["search_question_bank", "题库原题检索；仅在高置信且通过同题校验时返回权威解答"],
            ["retrieve_multi_rerank", "教材多路召回与精排；插图以可溯源占位符进入上下文"],
            ["graph_query", "查询知识点前置关系、关联链路与相关练习线索"],
            ["read_profile", "读取学情聚合画像，供个性化讲解与练习决策"],
            ["memory_search / upsert / forget", "长期记忆的检索、写入与按需遗忘"],
        ],
        widths=[4.5, 10.7],
    )

    h(doc, "5.4 知识底座", 2)
    body(
        doc,
        "知识工程采用「离线灌库 / 在线检索」解耦：教材按章切块并向量化，题库同时进入业务表与向量集合，"
        "知识图谱以校本结构化数据提供路径推理。"
        "在线路径只消费索引与元数据，不在请求链路中现场解析整本教材，从而保证时延与可重复性。",
    )

    h(doc, "5.5 产品应用构件", 2)
    table(
        doc,
        ["应用面", "与中枢关系", "主要能力"],
        [
            ["自由问答", "智能体主入口", "多模态感知、短/长路径编排、流式讲解与出图"],
            ["自适应练习", "消费学情与题库", "按薄弱点抽题、即时反馈、错题沉淀"],
            ["错题与学情", "中枢输出的沉淀层", "错题管理、画像批处理、学习反馈"],
            ["知识图谱页", "知识底座的可视投影", "概念关系与路径浏览"],
            ["管理控制台", "内容经营与运维面", "题库 CRUD、回流审核、班级学情、执行轨迹"],
        ],
        widths=[3.2, 4.0, 8.0],
    )

    # —— 6 运行时 ——
    h(doc, "6. 运行时视图", 1)
    h(doc, "6.1 端到端控制流", 2)
    body(
        doc,
        "一次学习请求首先进入感知与意图路由。路由模型仅输出结构化路由结果（任务类型、执行模式、原子任务清单、安全裁决），"
        "不承担解题与绘图。"
        "短路径请求直接进入解题编排；长路径请求先生成任务清单，经对话内人在回路确认后，按原子任务串行进入同一解题编排。"
        "单回合原子任务规模设有上限，超额部分通过续跑机制推进，以控制单次上下文与工具预算。",
    )

    h(doc, "6.2 意图双模式", 2)
    table(
        doc,
        ["模式", "适用情境", "控制策略"],
        [
            ["react_short", "寒暄、单点讲解、单题及少量原子任务", "跳过强制确认，直接进入解题编排与工具循环"],
            ["plan_execute", "多题、整卷、长文档、大纲类复合任务", "任务清单确认后方可执行；支持中断后续跑"],
        ],
        widths=[3.2, 5.8, 6.2],
    )
    body(
        doc,
        "安全策略将明显越界或违规请求标记为阻断并返回说明文案，使其不得进入工具执行链。"
        "路由失败时降级为「整段单一任务」，保证可用性优先于过度拆分。",
    )

    h(doc, "6.3 解题编排图", 2)
    body(
        doc,
        "解题过程建模为有状态控制图。节点职责单一：归一与识读、轻量决策、题库预取、知识锚定、工具增强求解、结果校验与回流写入。"
        "题库预取构成权威短路：命中且通过校验则直接以权威语义交付；未命中则进入知识锚定与 ReAct 工具循环，再经校验决定是否进入回流队列。",
    )
    code(
        doc,
        """normalize → ocr_or_text → decide → scope_task → bank_prefetch
        │                                        │
        │                              命中 ──► bank_done ──► 权威交付
        │                                        │
        │                              未命中 ─► kg_ground → ai_solve
        │                                                      │
        └──────────────────────────────────────────────► validate
                                                              │
                                                    回流写入 / 跳过回流""",
    )
    body(
        doc,
        "决策节点输出意图与步骤标记，不直接生成最终答卷；真正的工具选择发生在求解节点的 ReAct 循环中。"
        "交付层区分权威解答与 AI 参考解答，避免信任语义混淆。适合沉淀的未命中题目进入回流，由管理端审核后入库，形成内容飞轮。",
    )

    h(doc, "6.4 导师内核组装", 2)
    body(
        doc,
        "开放式求解内核由聊天模型、系统提示（学科人设、场景边界、题库约定、工具清单与绘图规约）、"
        "工具注册表与图执行器共同组成。"
        "运行时注入题面、预检索说明、记忆摘要与图谱上下文等原始材料，提示模板在节点内按需装配，"
        "符合「状态保存原始数据、提示在边界处格式化」的编排原则。",
    )
    table(
        doc,
        ["控制维度", "策略"],
        [
            ["迭代深度", "限制 ReAct 递归步数，防止无界自我循环"],
            ["时间预算", "墙钟超时与空闲超时共同约束"],
            ["工具配额", "绘图、教材检索、题库搜索分别计数并熔断"],
            ["重复动作", "同名同参连续重复触发断路"],
            ["输出规模", "限制最大输出长度，抑制生成坍缩"],
        ],
        widths=[3.5, 11.7],
    )

    h(doc, "6.5 绘图工作流", 2)
    body(
        doc,
        "专业出图被设计为解题工作流的附属能力，而非独立设计工具产品。"
        "调用方仅面对统一门面：选择或推断图类，生成结构化中间表示，经模式校验与有限重试后，交由确定性渲染引擎输出 SVG。"
        "答案正文使用占位标记绑定真实产物地址，禁止模型虚构媒体链接。",
    )
    table(
        doc,
        ["图类", "典型用途", "渲染路径"],
        [
            ["logic_dag", "门级逻辑与表达式实现", "逻辑 IR → netlistsvg"],
            ["timing_wave", "多信号时序波形", "波形 IR → 专用渲染"],
            ["state_machine", "状态转换图", "状态 IR → 专用渲染"],
            ["truth_table / kmap / seven_seg", "真值表、卡诺图、七段显示", "参数化 Python 绘制"],
            ["char_curve", "TTL/CMOS 传输特性", "科学绘图库"],
            ["msi_design", "74/555/CMOS/DAC 等设计电路", "逻辑先行校验 → 电路渲染引擎"],
        ],
        widths=[4.2, 5.0, 6.0],
    )
    body(
        doc,
        "设计电路类题目强制「逻辑先行」：先形成可校验的逻辑方案，再布局出图，扩展点落在器件引脚库而非按题定制模板。"
        "会话级绘图调用另有配额；若题面明确要求出图而产物缺失，求解收尾阶段可再次进入同一工作流补全，而不引入旁路绘制接口。",
    )

    h(doc, "6.6 记忆与个性化", 2)
    table(
        doc,
        ["层次", "作用"],
        [
            ["短期会话", "服务端会话消息作为权威历史，支撑多轮连贯"],
            ["长期记忆", "偏好、易错点与笔记经工具读写，跨会话复用"],
            ["学情画像", "掌握度与薄弱点聚合，服务练习选题与讲解侧重"],
        ],
        widths=[3.2, 12.0],
    )

    # —— 7 部署 ——
    h(doc, "7. 部署视图", 1)
    body(
        doc,
        "标准演示环境采用容器编排，将前端静态资源、API 服务与配套数据组件一并拉起。"
        "资源受限场景可采用精简拓扑：反向代理 + 应用进程 + 嵌入式/轻量数据库，在保持主链路可用的前提下降低运维面。"
        "绘图能力依赖对应运行时（脚本引擎、工作台资源与设计电路渲染环境）；部署说明中单独给出环境矩阵与自检入口。",
    )

    # —— 8 架构决策 ——
    h(doc, "8. 关键架构决策", 1)
    table(
        doc,
        ["决策", "选择", "后果"],
        [
            ["编排骨架", "LangGraph 状态图 + ReAct 工具循环", "分支、短路与循环显式化；便于插入校验与回流"],
            ["任务形态", "短路径 / 长路径双模式 + 人在回路", "控制长任务失控，同时保留单题低摩擦"],
            ["权威策略", "题库预取短路 + 交付分轨", "提升可信度；未命中进入回流经营"],
            ["绘图策略", "单一工具门面 + IR 校验 + 确定性渲染", "降低选错工具概率；图质量可回归"],
            ["知识策略", "离线灌库与在线检索分离", "时延稳定；教材更新走灌库流水线"],
            ["角色策略", "助学中枢 + 管理经营面", "先立助学深度；向教考协同平滑扩展"],
        ],
        widths=[3.0, 5.5, 6.7],
    )

    # —— 9 展望 ——
    h(doc, "9. 未来迭代与产品展望", 1)
    body(
        doc,
        "现行架构已形成可运行的助学中枢与内容经营闭环。在此基础上，产品将沿「学得更深、教得更省、管得更清」三条主线持续演进："
        "把自由探究式答疑延伸为有进度的课程护航，把题库经营延伸为教考协同的多智能体协作，并把校本适配与规模化验证做成可复制的学科样板。"
        "下列方向描述目标形态与协作关系，作为版本路线图的展望，而非对现状的否定。",
    )

    h(doc, "9.1 多智能体教考协同", 2)
    body(
        doc,
        "展望在助学智能体旁侧生长考务智能体簇，形成「学」与「考」分工协作的多 Agent 拓扑："
        "组卷智能体依据大纲与题型约束生成试卷结构；考务编排贯通发布、作答会话与收卷；"
        "判卷智能体完成客观题自动判定，并为主观题提供初评以辅助教师改判；洞察服务从考次维度沉淀失分结构与共性错因。"
        "协作上采用中心编排（Orchestrator）与专业执行者（Worker）相结合，助学与考务分工具域、分权限与分审计轨迹，"
        "使成绩相关链路具备可追溯治理能力，同时复用已有题库、学情与可观测基础设施。",
    )

    h(doc, "9.2 课程制护航学习（Curriculum Agent）", 2)
    body(
        doc,
        "在自由问答之外，进一步建设按章推进的课程制教学智能体："
        "以课程状态机维护「当前节次、学习目标、待练题与掌握度」；课内步骤按需调用教材检索与练习工具；"
        "互动上结合苏格拉底式提问与掌握度门禁，章末生成学情报告并驱动加练。"
        "届时，自由问答继续承担探究式自习，课程智能体承担有进度的护航教学，二者共享知识底座与画像，形成「能问、能学、能巩固」的完整体验。",
    )

    h(doc, "9.3 平台能力持续增强", 2)
    bullet(doc, "支持校本大纲裁剪与章节进度配置，服务不同院校的开课差异。", lead="教学配置：")
    bullet(doc, "图谱查询与编排状态进一步工程化，配套学科评测集持续校准效果。", lead="知识与编排：")
    bullet(doc, "对高频窄域图类引入专用小模型路径，与主工作流形成成本—质量协同。", lead="绘图效能：")
    bullet(doc, "扩大班级规模验证与学习效果度量，沉淀可复制的学科落地样板。", lead="推广验证：")

    h(doc, "9.4 版本路线展望", 2)
    code(
        doc,
        """① 持续打磨助学中枢体验、评测与绘图运行时
② 推出 Curriculum Agent 单章护航学习闭环
③ 建设组卷规则与题库标签，开展教考协同试点
④ 完善助学/考务多智能体协作与统一可观测
⑤ 开放校本配置能力，推进多班级规模化应用""",
    )

    # —— 10 小结 ——
    h(doc, "10. 小结", 1)
    body(
        doc,
        "学灵伴以助学智能体为架构中枢，用分层运行时、双模式任务控制、权威优先的解题编排、"
        "收敛的工具注册表与教材级绘图工作流，把学科垂类大模型能力落实为可部署的教学生产系统；"
        "以管理控制台完成内容经营与可观测闭环。"
        "面向未来，课程制护航学习与多智能体教考协同构成清晰的产品展望，使系统能够在现有中枢之上平滑生长，"
        "持续服务「人工智能 + 一流学科」场景。",
    )

    p = doc.add_paragraph()
    fmt(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=0)
    set_run_font(p.add_run("— 全文完 —"), east_asia="宋体", size_pt=10.5, color=RGBColor(0x66, 0x66, 0x66))

    doc.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    build()
